"""파이프라인 통합 테스트 및 슬라이드 범위 검증"""

import io
import pytest
from pptx import Presentation
from pptx.util import Inches
from PIL import Image

from src.main import parse_slide_ranges, load_config
from src.analyzer import SlideAnalyzer
from src.classifier import ElementClassifier
from src.composer import SlideComposer
from src.builder import NativePPTXBuilder
from src.diagram_analyzer import DiagramData


class TestParseSlideRanges:
    """슬라이드 범위 파싱 테스트"""

    def test_single_slide(self):
        assert parse_slide_ranges('3') == [2]

    def test_multiple_slides(self):
        assert parse_slide_ranges('1,3,5') == [0, 2, 4]

    def test_range(self):
        assert parse_slide_ranges('2-5') == [1, 2, 3, 4]

    def test_mixed(self):
        assert parse_slide_ranges('1,3-5,8') == [0, 2, 3, 4, 7]

    def test_duplicates_removed(self):
        assert parse_slide_ranges('1,1,2,2') == [0, 1]

    def test_max_slide_filter(self):
        """최대 슬라이드를 초과하는 번호가 필터링되는지 확인"""
        result = parse_slide_ranges('1,5,10,20', max_slide=10)
        assert result == [0, 4, 9]

    def test_max_slide_range_filter(self):
        """범위 지정 시 최대 슬라이드를 초과하는 부분이 제거되는지 확인"""
        result = parse_slide_ranges('8-15', max_slide=10)
        assert result == [7, 8, 9]

    def test_all_exceed_max(self):
        """모든 번호가 최대를 초과하면 빈 리스트 반환"""
        result = parse_slide_ranges('20,30', max_slide=10)
        assert result == []

    def test_invalid_format_ignored(self):
        """잘못된 형식은 무시"""
        result = parse_slide_ranges('1,abc,3')
        assert result == [0, 2]

    def test_empty_parts_ignored(self):
        result = parse_slide_ranges('1,,3,')
        assert result == [0, 2]

    def test_zero_and_negative_ignored(self):
        result = parse_slide_ranges('0,-1,1,2')
        assert result == [0, 1]


class TestComposerIntegration:
    """SlideComposer 통합 테스트"""

    @pytest.fixture
    def sample_pptx(self, tmp_path):
        """테스트용 PPTX 파일"""
        pptx_path = tmp_path / 'test.pptx'
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        layout = prs.slide_layouts[5]
        slide = prs.slides.add_slide(layout)
        txbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        txbox.text_frame.text = '테스트 슬라이드'

        prs.save(str(pptx_path))
        return str(pptx_path)

    def test_no_rebuild_saves_copy(self, sample_pptx, tmp_path):
        """재구성 필요 없으면 원본 그대로 저장"""
        output_path = str(tmp_path / 'output.pptx')

        analyzer = SlideAnalyzer()
        analyses = analyzer.analyze(sample_pptx)
        classifier = ElementClassifier()
        classifications = classifier.classify_all(analyses)

        composer = SlideComposer()
        result = composer.compose(
            sample_pptx, analyses, classifications, {}, output_path,
        )

        assert result == output_path
        prs = Presentation(output_path)
        assert len(prs.slides) == 1


class TestLoadConfig:
    """설정 로드 테스트"""

    def test_load_default_config(self):
        """기본 config.yaml 로드"""
        config = load_config()
        assert 'analysis' in config
        assert 'builder' in config
        assert 'output' in config

    def test_load_nonexistent_config(self):
        """존재하지 않는 설정 파일은 빈 딕셔너리 반환"""
        config = load_config('/nonexistent/config.yaml')
        assert config == {}

    def test_config_has_new_fields(self):
        """새로 추가된 설정 필드 확인"""
        config = load_config()
        analysis = config.get('analysis', {})
        assert 'api_timeout' in analysis
        assert 'max_retries' in analysis


class TestBuilderEdgeCases:
    """빌더 엣지 케이스 테스트"""

    @pytest.fixture
    def builder(self):
        return NativePPTXBuilder()

    @pytest.fixture
    def blank_slide(self):
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        return prs, slide

    def test_build_unknown_type_falls_back(self, builder, blank_slide):
        """알 수 없는 diagram_type은 process로 폴백"""
        prs, slide = blank_slide
        data = DiagramData(
            diagram_type='some_future_type',
            nodes=[
                {'id': 'n1', 'text': 'Test', 'color_fill': '#4A90D9', 'color_text': '#FFFFFF'},
            ],
        )
        builder.build(slide, data)
        assert len(list(slide.shapes)) > 0

    def test_empty_diagram_data(self, builder, blank_slide):
        """빈 DiagramData로 빌드 시 에러 없음"""
        prs, slide = blank_slide
        data = DiagramData()
        builder.build(slide, data)
        assert len(list(slide.shapes)) == 0

    def test_node_with_none_text(self, builder, blank_slide):
        """text가 None인 노드 처리"""
        prs, slide = blank_slide
        data = DiagramData(
            diagram_type='process',
            nodes=[
                {'id': 'n1', 'text': None, 'color_fill': '#4A90D9', 'color_text': '#FFFFFF'},
            ],
        )
        builder.build_process_diagram(slide, data)
        assert len(list(slide.shapes)) > 0
