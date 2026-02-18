"""Flask API 테스트"""

from __future__ import annotations

import io
import pytest
from pptx import Presentation

from src.api import app


@pytest.fixture
def client():
    """Flask 테스트 클라이언트"""
    app.config['TESTING'] = True
    with app.test_client() as c:
        yield c


@pytest.fixture
def sample_pptx_bytes():
    """테스트용 PPTX 바이너리 생성"""
    prs = Presentation()

    # 텍스트 슬라이드
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    from pptx.util import Inches
    txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    txBox.text_frame.text = "테스트 슬라이드"

    # 테이블 슬라이드
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    table = slide2.shapes.add_table(3, 3, Inches(1), Inches(1), Inches(5), Inches(3)).table
    table.cell(0, 0).text = "헤더1"
    table.cell(0, 1).text = "헤더2"

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf.getvalue()


def test_health(client):
    """헬스체크 엔드포인트 테스트"""
    resp = client.get('/api/health')
    assert resp.status_code == 200
    assert resp.json['status'] == 'ok'


def test_analyze_no_file(client):
    """파일 없이 분석 요청 시 400 반환"""
    resp = client.post('/api/analyze')
    assert resp.status_code == 400
    assert 'error' in resp.json


def test_analyze_invalid_file(client):
    """잘못된 파일 형식 시 400 반환"""
    data = {'file': (io.BytesIO(b'not a pptx'), 'test.txt')}
    resp = client.post('/api/analyze', data=data, content_type='multipart/form-data')
    assert resp.status_code == 400


def test_analyze_success(client, sample_pptx_bytes):
    """정상 분석 요청 테스트"""
    data = {'file': (io.BytesIO(sample_pptx_bytes), 'test.pptx')}
    resp = client.post('/api/analyze', data=data, content_type='multipart/form-data')
    assert resp.status_code == 200
    result = resp.json
    assert 'file_id' in result
    assert result['total_slides'] == 2
    assert isinstance(result['slides'], list)
    assert len(result['slides']) == 2


def test_analyze_with_slides_filter(client, sample_pptx_bytes):
    """슬라이드 필터링 테스트"""
    data = {
        'file': (io.BytesIO(sample_pptx_bytes), 'test.pptx'),
        'slides': '1',
    }
    resp = client.post('/api/analyze', data=data, content_type='multipart/form-data')
    assert resp.status_code == 200
    assert resp.json['total_slides'] == 1


def test_process_no_file(client):
    """파일 없이 변환 요청 시 400 반환"""
    resp = client.post('/api/process')
    assert resp.status_code == 400


def test_process_success(client, sample_pptx_bytes):
    """정상 변환 요청 테스트 (도식 없는 파일)"""
    data = {'file': (io.BytesIO(sample_pptx_bytes), 'test.pptx')}
    resp = client.post('/api/process', data=data, content_type='multipart/form-data')
    assert resp.status_code == 200
    result = resp.json
    # 도식이 없으므로 download_url은 None
    assert result['download_url'] is None


def test_index_returns_html(client):
    """루트 경로가 Task Pane HTML을 반환하는지 확인"""
    resp = client.get('/')
    assert resp.status_code == 200
